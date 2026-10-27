from pathlib import Path
import io

import pandas as pd
import plotly.express as px
import streamlit as st


# =====================================================
# 1. CONFIGURACIÓN
# =====================================================

st.set_page_config(
    page_title="Visualizador de Prestaciones",
    page_icon="📊",
    layout="wide",
    initial_sidebar_state="expanded"
)

st.markdown("""
<style>
.block-container {
    padding-top: 1rem;
    padding-bottom: 1rem;
    max-width: 100%;
}

h1 { font-size: 1.5rem !important; }
h2 { font-size: 1.2rem !important; }
h3 { font-size: 1rem !important; }

button[data-baseweb="tab"] {
    font-size: 0.85rem;
    padding: 8px 12px;
}
</style>
""", unsafe_allow_html=True)


# =====================================================
# 2. CONSTANTES
# =====================================================

REQUIRED_COLUMNS = [
    "CODIGO_PRESTACION",
    "DESCRIPCION_PRESTACION",
    "FECHA_INICIO_VIGENCIA",
    "FECHA_FIN_VIGENCIA",
    "VALOR"
]

COMPARE_COLUMNS = [
    "CATEGORIA",
    "NNE",
    "DESCRIPCION",
    "VALOR"
]

DEFAULT_FILE_PATH = Path(__file__).with_name("valores.xlsx")
DEFAULT_COMPARE_PATH = Path(__file__).with_name("administracion.xlsx")
DEFAULT_ALFABETA_PATH = Path(__file__).with_name("alfabeta.xlsx")


# =====================================================
# 3. FUNCIONES AUXILIARES
# =====================================================

@st.cache_data(show_spinner=True)
def load_excel(file_bytes):
    return pd.read_excel(
        io.BytesIO(file_bytes),
        engine="openpyxl"
    )


def period_label(x):
    """Convierte 2026-01 a Ene-2026."""

    meses = [
        "Ene", "Feb", "Mar", "Abr",
        "May", "Jun", "Jul", "Ago",
        "Sep", "Oct", "Nov", "Dic"
    ]

    fecha = pd.Timestamp(x)

    return f"{meses[fecha.month - 1]}-{fecha.year}"


def format_money(v):
    """Formato monetario argentino."""

    if pd.isna(v):
        return "-"

    return "$" + (
        f"{v:,.2f}"
        .replace(",", "X")
        .replace(".", ",")
        .replace("X", ".")
    )


def normalize_value(series):
    """Convierte valores numéricos a formato decimal."""

    return pd.to_numeric(series, errors="coerce")


# =====================================================
# 3b. ALERTAS, SEMÁFORO Y EXPORTACIÓN
# =====================================================

def to_excel_bytes(sheets):
    """sheets = {"Nombre hoja": dataframe, ...} -> bytes de un .xlsx"""
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        for name, d in sheets.items():
            d.to_excel(writer, sheet_name=name[:31], index=False)
    return buf.getvalue()


def build_alerts(historical_df, monthly_df, umbral_suba, umbral_baja):
    """Alertas de calidad sobre el histórico."""

    cols = ["CODIGO_PRESTACION", "DESCRIPCION_PRESTACION",
            "PERIODO", "VALOR", "VARIACION_PCT"]
    partes = []

    suba = monthly_df[monthly_df["VARIACION_PCT"] > umbral_suba][cols].copy()
    suba["ALERTA"] = f"Suba mayor a {umbral_suba:g}%"
    partes.append(suba)

    baja = monthly_df[monthly_df["VARIACION_PCT"] < -umbral_baja][cols].copy()
    baja["ALERTA"] = (
        "Baja (variación negativa)" if umbral_baja == 0
        else f"Baja mayor a {umbral_baja:g}%"
    )
    partes.append(baja)

    cero = historical_df[historical_df["VALOR"] <= 0].copy()
    cero["VARIACION_PCT"] = float("nan")
    cero = cero[[c for c in cols]].copy()
    cero["ALERTA"] = "Valor cero o negativo"
    partes.append(cero)

    # Repetidos en un mismo mes (se usa el histórico completo,
    # porque la versión mensual ya se quedó con el último)
    g = historical_df.groupby(["CODIGO_PRESTACION", "PERIODO"])["VALOR"]
    rep = g.agg(["count", "nunique"]).reset_index()
    rep = rep[rep["count"] > 1]
    if not rep.empty:
        rep["ALERTA"] = rep["nunique"].map(
            lambda n: "Repetido en el mes con valores distintos"
            if n > 1 else "Repetido en el mes (mismo valor)"
        )
        rep = rep.merge(
            monthly_df[cols], on=["CODIGO_PRESTACION", "PERIODO"], how="left"
        )[cols + ["ALERTA"]]
        partes.append(rep)

    out = pd.concat(partes, ignore_index=True)
    return out.sort_values(["ALERTA", "CODIGO_PRESTACION", "PERIODO"])


def semaforo(pct, tol_verde, tol_amarillo):
    if pd.isna(pct):
        return "⚪ Sin dato"
    a = abs(pct)
    if a <= tol_verde:
        return "🟢 OK"
    if a <= tol_amarillo:
        return "🟡 Revisar"
    return "🔴 Desvío"


def style_workbook(writer, money_cols, pct_cols):
    """Formato: encabezado, filtros, panel fijo, anchos, números y semáforo."""
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    head_fill = PatternFill("solid", fgColor="1F3864")
    fills = {
        "🟢": PatternFill("solid", fgColor="C6EFCE"),
        "🟡": PatternFill("solid", fgColor="FFEB9C"),
        "🔴": PatternFill("solid", fgColor="FFC7CE"),
        "⚪": PatternFill("solid", fgColor="E7E6E6"),
    }

    for ws in writer.book.worksheets:
        if ws.max_row < 1 or ws.max_column < 1:
            continue

        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = head_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        ws.auto_filter.ref = ws.dimensions
        ws.freeze_panes = "C2" if ws.title != "Resumen" else None

        if ws.title == "Resumen":
            ws.column_dimensions["A"].width = 62
            ws.column_dimensions["B"].width = 55
            for row in ws.iter_rows(min_row=2):
                if row[0].value and not str(row[1].value or "").strip() \
                        and str(row[0].value).isupper():
                    row[0].font = Font(bold=True)
            continue

        for idx, cell in enumerate(ws[1], start=1):
            name = str(cell.value)
            letter = get_column_letter(idx)

            if name in ("DESCRIPCION", "Descripción"):
                ws.column_dimensions[letter].width = 60
            elif name == "CODIGO_PRESTACION":
                ws.column_dimensions[letter].width = 16
            elif name in ("SEMAFORO", "ESTADO", "COINCIDE_CON"):
                ws.column_dimensions[letter].width = 30
            else:
                ws.column_dimensions[letter].width = 16

            fmt = None
            if name in money_cols:
                fmt = '"$"#,##0.00'
            elif name in pct_cols:
                fmt = '0.00"%"'

            for row in ws.iter_rows(min_row=2, min_col=idx, max_col=idx):
                c = row[0]
                if fmt:
                    c.number_format = fmt
                if name == "SEMAFORO" and c.value:
                    f = fills.get(str(c.value)[:1])
                    if f:
                        c.fill = f


def build_report(t, df, compare_period, verde, amarillo, tol=0.1, top_n=100,
                 alfabeta=None, nombres=None):
    """Informe Excel completo del listado de administración."""

    base_cols = ["SEMAFORO", "COINCIDE_CON", "ESTADO",
                 "CODIGO_PRESTACION", "DESCRIPCION",
                 "VALOR_ADMINISTRACION", "VALOR_ALFABETA",
                 "DIF_ALFABETA_VS_ADMIN", "DIF_ALFABETA_VS_ADMIN_PCT",
                 "VALOR_HISTORICO", "PERIODO_HISTORICO", "VALORIZADO_EN_PERIODO",
                 "DIFERENCIA", "DIFERENCIA_PCT",
                 "VALOR_HIST_N1", "PERIODO_N1",
                 "DIFERENCIA_N1", "DIFERENCIA_PCT_N1",
                 "VAR_HIST_N_VS_N1_PCT", "N_VALORIZACIONES"]

    # Hoja 1: listado de administración vs último valor histórico
    hoja_adm = t[base_cols].copy()

    # Hoja 2: movimiento histórico de cada código de administración
    periodos = [p for p in sorted(df["PERIODO"].unique()) if p <= compare_period]
    etiquetas = {p: period_label(p) for p in periodos}

    mov = (
        df[df["CODIGO_PRESTACION"].isin(t["CODIGO_PRESTACION"])
           & (df["PERIODO"] <= compare_period)]
        .pivot_table(index="CODIGO_PRESTACION", columns="PERIODO",
                     values="VALOR", aggfunc="last")
        .reindex(columns=periodos)
        .rename(columns=etiquetas)
    )

    hoja_mov = t[["CODIGO_PRESTACION", "DESCRIPCION"]].merge(
        mov, left_on="CODIGO_PRESTACION", right_index=True, how="left"
    )
    hoja_mov["ADMINISTRACION"] = t["VALOR_ADMINISTRACION"].values
    hoja_mov["ALFABETA"] = t["VALOR_ALFABETA"].values
    hoja_mov["DIF_ALFABETA_VS_ADMIN_PCT"] = t["DIF_ALFABETA_VS_ADMIN_PCT"].values
    hoja_mov["DIFERENCIA_PCT"] = t["DIFERENCIA_PCT"].values
    hoja_mov["DIFERENCIA_PCT_N1"] = t["DIFERENCIA_PCT_N1"].values
    hoja_mov["COINCIDE_CON"] = t["COINCIDE_CON"].values

    # Hoja 3: mayores desvíos
    con_dato = t.dropna(subset=["DIFERENCIA_PCT"]).copy()
    desvios = con_dato.reindex(
        con_dato["DIFERENCIA_PCT"].abs().sort_values(ascending=False).index
    ).head(top_n)[base_cols]

    # Hoja 4: sin valor histórico
    sin_dato = t[t["VALOR_HISTORICO"].isna()][
        ["ESTADO", "CODIGO_PRESTACION", "DESCRIPCION", "VALOR_ADMINISTRACION",
         "VALOR_ALFABETA"]
    ]

    # Hoja 0: resumen
    cuenta = t["SEMAFORO"].value_counts()
    nombres = nombres or {}
    alf = None
    if alfabeta is not None:
        alf = alfabeta.drop_duplicates("CODIGO_PRESTACION", keep="last")[
            ["CODIGO_PRESTACION", "DESCRIPCION", "VALOR"]
        ].rename(columns={"VALOR": "VALOR_ALFABETA", "DESCRIPCION": "DESC_ALFABETA"})

    hasta_res = df[df["PERIODO"] <= compare_period]
    filas = [
        ("ORÍGENES", ""),
        ("Prestaciones (histórico): archivo", nombres.get("historico", "-")),
        ("Prestaciones (histórico): códigos distintos", int(df["CODIGO_PRESTACION"].nunique())),
        ("Prestaciones (histórico): códigos con valor hasta el período", int(hasta_res["CODIGO_PRESTACION"].nunique())),
        ("Administración: archivo", nombres.get("administracion", "-")),
        ("Administración: códigos", len(t)),
        ("Alfabeta: archivo", nombres.get("alfabeta", "-") if alf is not None else "sin archivo"),
    ]
    if alf is not None:
        en_adm = alf["CODIGO_PRESTACION"].isin(t["CODIGO_PRESTACION"])
        con_alf = t["VALOR_ALFABETA"].notna()
        coinc = (t["DIF_ALFABETA_VS_ADMIN_PCT"].abs() <= tol).sum()
        filas += [
            ("Alfabeta: códigos", len(alf)),
            ("Alfabeta: códigos que también están en administración", int(en_adm.sum())),
            ("Alfabeta: códigos que no están en administración", int((~en_adm).sum())),
            ("Administración sin precio en Alfabeta", int((~con_alf).sum())),
            ("Alfabeta igual a administración (dentro de tolerancia)", int(coinc)),
            ("Alfabeta distinto de administración", int(con_alf.sum() - coinc)),
        ]
    filas += [
        ("", ""),
        ("PARÁMETROS", ""),
        ("Período de comparación", period_label(compare_period)),
        ("Verde hasta ±%", verde),
        ("Amarillo hasta ±%", amarillo),
        ("Coincidencia: tolerancia ±%", tol),
        ("", ""),
        ("SEMÁFORO ADMINISTRACIÓN VS HISTÓRICO", ""),
    ] + [(k, int(v)) for k, v in cuenta.items()] + [
        (f"Administración vs histórico: {k}", int(v))
        for k, v in t["COINCIDE_CON"].value_counts().items()
    ] + [
        ("N", "Última valorización histórica hasta el período elegido"),
        ("N-1", "Última valorización histórica anterior al período elegido"),
    ]
    resumen = pd.DataFrame(filas, columns=["Concepto", "Valor"])

    hoja_alf = None
    if alf is not None:
        vig = (
            df[df["PERIODO"] <= compare_period].sort_values("PERIODO")
            .drop_duplicates("CODIGO_PRESTACION", keep="last")
            [["CODIGO_PRESTACION", "PERIODO", "VALOR"]]
            .rename(columns={"PERIODO": "PERIODO_HISTORICO",
                             "VALOR": "VALOR_HISTORICO"})
        )
        adm_v = t[["CODIGO_PRESTACION", "DESCRIPCION", "VALOR_ADMINISTRACION"]]
        hoja_alf = alf.merge(adm_v, on="CODIGO_PRESTACION", how="left")
        hoja_alf = hoja_alf.merge(vig, on="CODIGO_PRESTACION", how="left")
        hoja_alf["DESCRIPCION"] = hoja_alf["DESCRIPCION"].fillna(hoja_alf["DESC_ALFABETA"])
        hoja_alf["ESTADO"] = hoja_alf["VALOR_ADMINISTRACION"].notna().map(
            {True: "En administración", False: "Solo en Alfabeta"}
        )
        hoja_alf["DIF_VS_ADMIN"] = hoja_alf["VALOR_ALFABETA"] - hoja_alf["VALOR_ADMINISTRACION"]
        hoja_alf["DIF_VS_ADMIN_PCT"] = (
            hoja_alf["DIF_VS_ADMIN"] / hoja_alf["VALOR_ADMINISTRACION"].replace(0, float("nan"))
        ) * 100
        hoja_alf["DIF_VS_HISTORICO"] = hoja_alf["VALOR_ALFABETA"] - hoja_alf["VALOR_HISTORICO"]
        hoja_alf["DIF_VS_HISTORICO_PCT"] = (
            hoja_alf["DIF_VS_HISTORICO"] / hoja_alf["VALOR_HISTORICO"].replace(0, float("nan"))
        ) * 100
        hoja_alf = hoja_alf[[
            "ESTADO", "CODIGO_PRESTACION", "DESCRIPCION",
            "VALOR_ALFABETA", "VALOR_ADMINISTRACION",
            "DIF_VS_ADMIN", "DIF_VS_ADMIN_PCT",
            "VALOR_HISTORICO", "PERIODO_HISTORICO",
            "DIF_VS_HISTORICO", "DIF_VS_HISTORICO_PCT",
        ]]

    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        resumen.to_excel(w, sheet_name="Resumen", index=False)
        hoja_adm.to_excel(w, sheet_name="Administracion", index=False)
        if hoja_alf is not None:
            hoja_alf.to_excel(w, sheet_name="Alfabeta", index=False)
        hoja_mov.to_excel(w, sheet_name="Historico vs Admin", index=False)
        desvios.to_excel(w, sheet_name="Mayores desvios", index=False)
        sin_dato.to_excel(w, sheet_name="Sin valor historico", index=False)

        style_workbook(
            w,
            money_cols={"VALOR_ADMINISTRACION", "VALOR_ALFABETA",
                        "DIF_ALFABETA_VS_ADMIN", "VALOR_HISTORICO", "DIFERENCIA",
                        "VALOR_HIST_N1", "DIFERENCIA_N1",
                        "ADMINISTRACION", "ALFABETA", "DIF_VS_ADMIN",
                        "DIF_VS_HISTORICO", *etiquetas.values()},
            pct_cols={"DIFERENCIA_PCT", "DIFERENCIA_PCT_N1",
                      "DIF_ALFABETA_VS_ADMIN_PCT", "DIF_VS_ADMIN_PCT",
                      "DIF_VS_HISTORICO_PCT",
                      "VAR_HIST_N_VS_N1_PCT"},
        )
    return buf.getvalue()


# =====================================================
# 4. LIMPIEZA DEL ARCHIVO HISTÓRICO
# =====================================================

def clean_data(raw):

    df = raw.copy()

    df.columns = [
        str(c).strip().upper()
        for c in df.columns
    ]

    missing = [
        c for c in REQUIRED_COLUMNS
        if c not in df.columns
    ]

    if missing:
        raise ValueError(
            f"Faltan columnas requeridas: {', '.join(missing)}"
        )

    df = df[REQUIRED_COLUMNS].copy()

    df["CODIGO_PRESTACION"] = (
        df["CODIGO_PRESTACION"]
        .astype("string")
        .str.strip()
    )

    df["DESCRIPCION_PRESTACION"] = (
        df["DESCRIPCION_PRESTACION"]
        .astype("string")
        .str.strip()
    )

    df["FECHA_INICIO_VIGENCIA"] = pd.to_datetime(
        df["FECHA_INICIO_VIGENCIA"],
        errors="coerce"
    )

    df["FECHA_FIN_VIGENCIA"] = pd.to_datetime(
        df["FECHA_FIN_VIGENCIA"],
        errors="coerce"
    )

    df["VALOR"] = normalize_value(df["VALOR"])

    # Eliminar registros sin datos fundamentales
    df = df.dropna(
        subset=[
            "CODIGO_PRESTACION",
            "FECHA_INICIO_VIGENCIA",
            "VALOR"
        ]
    )

    df = df[df["CODIGO_PRESTACION"] != ""]

    # Generar período
    df["PERIODO"] = (
        df["FECHA_INICIO_VIGENCIA"]
        .dt.to_period("M")
        .astype(str)
    )

    # Ordenar
    df = df.sort_values([
        "FECHA_INICIO_VIGENCIA",
        "CODIGO_PRESTACION"
    ]).reset_index(drop=True)

    # Una valorización por código y período
    monthly = df.drop_duplicates(
        ["CODIGO_PRESTACION", "PERIODO"],
        keep="last"
    ).copy()

    monthly = monthly.sort_values([
        "CODIGO_PRESTACION",
        "PERIODO"
    ])

    # Variación contra la valorización anterior
    monthly["VARIACION_PCT"] = (
        monthly
        .groupby("CODIGO_PRESTACION")["VALOR"]
        .pct_change(fill_method=None) * 100
    )

    return df, monthly


# =====================================================
# 5. LIMPIEZA DEL SEGUNDO ARCHIVO
# =====================================================

def clean_compare_data(raw, period, origen="ADMINISTRACION"):

    df2 = raw.copy()

    df2.columns = [
        str(c).strip().upper()
        for c in df2.columns
    ]

    missing = [
        c for c in COMPARE_COLUMNS
        if c not in df2.columns
    ]

    if missing:
        raise ValueError(
            f"Faltan columnas: {', '.join(missing)}"
        )

    df2 = df2[COMPARE_COLUMNS].copy()

    # Convertir a texto ANTES de usar .str
    for col in ["CATEGORIA", "NNE", "DESCRIPCION"]:
        df2[col] = (
            df2[col]
            .astype("string")
            .str.strip()
        )

    # Si NNE vino como número, evitar que quede "101.0"
    df2["NNE"] = df2["NNE"].str.replace(
        r"\.0$", "", regex=True
    )

    df2["VALOR"] = normalize_value(df2["VALOR"])

    df2 = df2.dropna(
        subset=["CATEGORIA", "NNE", "VALOR"]
    )

    df2 = df2[
        (df2["CATEGORIA"] != "") & (df2["NNE"] != "")
    ]

    # Construir código de prestación
    df2["CODIGO_PRESTACION"] = (
        df2["CATEGORIA"].str.upper()
        + "."
        + df2["NNE"].str.upper()
    )

    df2["PERIODO"] = period

    df2["ORIGEN"] = origen

    return df2

# =====================================================
# 6. SIDEBAR - ARCHIVOS
# =====================================================

with st.sidebar:

    st.title("Configuración")

    with st.expander("1. Archivo histórico", expanded=True):

        upload_file = st.file_uploader(
            "Reemplazar archivo de valores (opcional)",
            type=["xlsx"],
            key="historico"
        )

    with st.expander("2. Archivo de comparación", expanded=True):

        compare_file = st.file_uploader(
            "Reemplazar planilla de administración (opcional)",
            type=["xlsx"],
            key="comparacion"
        )

    with st.expander("3. Archivo Alfabeta", expanded=True):

        alfabeta_file = st.file_uploader(
            "Reemplazar planilla de Alfabeta (opcional)",
            type=["xlsx"],
            key="alfabeta"
        )


# =====================================================
# 7. CARGAR ARCHIVO HISTÓRICO
# =====================================================

if upload_file is not None:

    file_bytes = upload_file.getvalue()
    source_name = upload_file.name

elif DEFAULT_FILE_PATH.exists():

    file_bytes = DEFAULT_FILE_PATH.read_bytes()
    source_name = DEFAULT_FILE_PATH.name

else:

    st.error(
        f"No se encontró el archivo: {DEFAULT_FILE_PATH.name}"
    )
    st.stop()


try:

    raw_data = load_excel(file_bytes)

    # Conservar todos los registros y versión mensual
    historical_df, df = clean_data(raw_data)

except Exception as e:

    st.error(f"Error procesando el archivo histórico: {e}")
    st.stop()


if df.empty:
    st.warning("El archivo histórico no contiene registros válidos.")
    st.stop()


# =====================================================
# 8. PERÍODOS
# =====================================================

periods = sorted(
    df["PERIODO"].dropna().unique()
)

period_names = {
    p: period_label(p)
    for p in periods
}

min_date = periods[0]
max_date = periods[-1]


# =====================================================
# 9. PERÍODO DEL SEGUNDO ARCHIVO
# =====================================================

with st.sidebar:

    compare_period = st.selectbox(
        "Período de comparación",
        options=periods,
        format_func=lambda p: period_names[p],
        index=len(periods) - 1
    )


# =====================================================
# 10. PROCESAR SEGUNDO ARCHIVO
# =====================================================

compare_df = None
compare_bytes = None
compare_name = None

if compare_file is not None:

    compare_bytes = compare_file.getvalue()
    compare_name = compare_file.name

elif DEFAULT_COMPARE_PATH.exists():

    compare_bytes = DEFAULT_COMPARE_PATH.read_bytes()
    compare_name = DEFAULT_COMPARE_PATH.name

alfabeta_df = None
alfabeta_bytes = None
alfabeta_name = None

if alfabeta_file is not None:

    alfabeta_bytes = alfabeta_file.getvalue()
    alfabeta_name = alfabeta_file.name

elif DEFAULT_ALFABETA_PATH.exists():

    alfabeta_bytes = DEFAULT_ALFABETA_PATH.read_bytes()
    alfabeta_name = DEFAULT_ALFABETA_PATH.name

with st.sidebar:

    st.caption(f"Histórico en uso: {source_name}")
    st.caption(
        f"Administración en uso: {compare_name}"
        if compare_name else "Administración: sin archivo"
    )
    st.caption(
        f"Alfabeta en uso: {alfabeta_name}"
        if alfabeta_name else "Alfabeta: sin archivo"
    )

if compare_bytes is not None:

    try:

        compare_raw = load_excel(compare_bytes)

        compare_df = clean_compare_data(
            compare_raw,
            compare_period
        )

    except Exception as e:

        st.error(
            f"Error procesando el archivo de comparación: {e}"
        )
        st.stop()


if alfabeta_bytes is not None:

    try:

        alfabeta_raw = load_excel(alfabeta_bytes)

        alfabeta_df = clean_compare_data(
            alfabeta_raw,
            compare_period,
            origen="ALFABETA"
        )

    except Exception as e:

        st.error(
            f"Error procesando el archivo de Alfabeta: {e}"
        )
        st.stop()


# =====================================================
# 11. ENCABEZADO Y MÉTRICAS
# =====================================================

st.title("Visualizador de prestaciones")

st.caption(
    f"Archivo histórico: {source_name} | "
    f"Administración: {compare_name or '-'} | "
    f"Alfabeta: {alfabeta_name or '-'} | "
    "Los meses sin nueva valorización quedan vacíos."
)

mostrar_metricas = st.toggle(
    "Mostrar resumen general",
    value=False
)

if mostrar_metricas:

    c1, c2, c3, c4 = st.columns(4)

    c1.metric(
        "Registros mensuales",
        f"{len(df):,}".replace(",", ".")
    )

    c2.metric(
        "Códigos únicos",
        f"{df['CODIGO_PRESTACION'].nunique():,}".replace(",", ".")
    )

    c3.metric(
        "Primer mes",
        period_label(min_date)
    )

    c4.metric(
        "Último mes",
        period_label(max_date)
    )


# =====================================================
# 12. METADATOS PARA FILTROS
# =====================================================

meta = (
    df.sort_values("FECHA_INICIO_VIGENCIA")
    .drop_duplicates("CODIGO_PRESTACION", keep="last")
    .set_index("CODIGO_PRESTACION")
)


# =====================================================
# 13. FILTROS
# =====================================================

st.subheader("Filtros")

f1, f2, f3 = st.columns([1.2, 2, 1.2])

search_code = f1.text_input(
    "Código contiene",
    placeholder="Ej. MED.X"
).strip().lower()

search_desc = f2.text_input(
    "Descripción contiene",
    placeholder="Ej. MORFINA"
).strip().lower()

period_filter = f3.selectbox(
    "Período",
    ["Todos"] + [period_names[p] for p in periods]
)


# Aplicar filtros por código y descripción

filtered_codes = meta.index.tolist()

if search_code:

    filtered_codes = [
        c for c in filtered_codes
        if search_code in str(c).lower()
    ]

if search_desc:

    filtered_codes = [
        c for c in filtered_codes
        if search_desc in str(
            meta.loc[c, "DESCRIPCION_PRESTACION"]
        ).lower()
    ]


# Datos filtrados por código

codes_df = df[
    df["CODIGO_PRESTACION"].isin(filtered_codes)
].copy()


# Aplicar período

filtered_df = codes_df.copy()

if period_filter != "Todos":

    chosen_period = next(
        p for p in periods
        if period_names[p] == period_filter
    )

    filtered_df = filtered_df[
        filtered_df["PERIODO"] == chosen_period
    ]


st.divider()


# =====================================================
# 14. CREACIÓN DE PESTAÑAS
# =====================================================

tab1, tab2, tab3, tab4, tab5 = st.tabs([

    "Valores",
    "Matriz histórica",
    "Variaciones",
    "Revisión por código",
    "Comparación de orígenes"

])


# =====================================================
# TAB 1 - VALORES
# =====================================================

with tab1:

    st.dataframe(
        filtered_df,
        use_container_width=True,
        hide_index=True,
        column_config={
            "VALOR": st.column_config.NumberColumn(
                "Valor",
                format="$ %.2f"
            ),
            "VARIACION_PCT": st.column_config.NumberColumn(
                "Variación %",
                format="%.2f %%"
            )
        }
    )

    st.caption(
        f"Registros encontrados: {len(filtered_df)}"
    )


# =====================================================
# TAB 2 - MATRIZ HISTÓRICA
# =====================================================

with tab2:

    matriz = filtered_df.pivot_table(

        index=[
            "CODIGO_PRESTACION",
            "DESCRIPCION_PRESTACION"
        ],

        columns="PERIODO",

        values="VALOR",

        aggfunc="last"
    )

    matriz = matriz.reindex(
        columns=[
            p for p in periods
            if p in matriz.columns
        ]
    )

    st.dataframe(
        matriz.style.format(format_money),
        use_container_width=True,
        height=600
    )

    st.divider()
    st.subheader("Alertas de calidad de datos")
    a1, a2 = st.columns(2)
    umbral_suba = a1.number_input("Alerta si sube más de (%)", value=100.0, step=10.0)
    umbral_baja = a2.number_input("Alerta si baja más de (%)", value=0.0, step=5.0,
                                  help="0 = cualquier variación negativa")
    alertas = build_alerts(
        historical_df[historical_df["CODIGO_PRESTACION"].isin(filtered_codes)],
        df[df["CODIGO_PRESTACION"].isin(filtered_codes)],
        umbral_suba, umbral_baja,
    )
    resumen = alertas["ALERTA"].value_counts()
    st.dataframe(resumen.rename("Cantidad"), use_container_width=True)
    tipos = st.multiselect("Tipo de alerta", sorted(alertas["ALERTA"].unique()),
                           default=sorted(alertas["ALERTA"].unique()))
    ver = alertas[alertas["ALERTA"].isin(tipos)]
    st.dataframe(
        ver, use_container_width=True, hide_index=True, height=400,
        column_config={
            "VALOR": st.column_config.NumberColumn("Valor", format="$ %.2f"),
            "VARIACION_PCT": st.column_config.NumberColumn("Variación %", format="%.2f %%"),
        },
    )
    st.download_button("⬇ Descargar alertas (Excel)",
                       to_excel_bytes({"Alertas": ver}),
                       file_name="alertas_calidad.xlsx")


# =====================================================
# TAB 3 - VARIACIONES
# =====================================================

with tab3:

    variacion = filtered_df.pivot_table(

        index="CODIGO_PRESTACION",

        columns="PERIODO",

        values="VARIACION_PCT",

        aggfunc="last"
    )

    st.dataframe(
        variacion.style.format(
            "{:.2f}%",
            na_rep="-"
        ),
        use_container_width=True,
        height=500
    )

    st.caption(
        "La variación se calcula contra la valorización "
        "anterior disponible de cada código."
    )


# =====================================================
# TAB 4 - REVISIÓN INDIVIDUAL
# =====================================================

with tab4:

    codigos = sorted(
        filtered_df["CODIGO_PRESTACION"].unique()
    )

    if codigos:

        codigo = st.selectbox(
            "Seleccionar código",
            codigos
        )

        # Utilizamos el histórico completo, sin eliminar
        # valorizaciones repetidas dentro de un mes.

        detalle = historical_df[
            historical_df["CODIGO_PRESTACION"] == codigo
        ].copy()

        detalle = detalle.sort_values(
            "FECHA_INICIO_VIGENCIA"
        )

        st.dataframe(
            detalle,
            use_container_width=True,
            hide_index=True
        )

        fig = px.line(

            detalle,

            x="FECHA_INICIO_VIGENCIA",

            y="VALOR",

            markers=True,

            title=f"Evolución histórica - {codigo}"
        )

        fig.update_layout(
            xaxis_title="Fecha de vigencia",
            yaxis_title="Valor",
            height=420
        )

        st.plotly_chart(
            fig,
            use_container_width=True
        )

    else:

        st.info(
            "No hay códigos disponibles con los filtros actuales."
        )


# =====================================================
# TAB 5 - COMPARACIÓN DE ORÍGENES
# =====================================================

with tab5:

    if compare_df is None:
        st.info(
            "No hay planilla de administración. Subila desde el panel "
            "lateral o dejá administracion.xlsx junto a la app."
        )

    else:
        adm = compare_df.copy()

        # Los filtros superiores (código / descripción del histórico)
        # solo se aplican si el usuario los está usando.
        if search_code or search_desc:
            adm = adm[adm["CODIGO_PRESTACION"].isin(filtered_codes)]

        # Duplicados en administración: se avisa y se conserva el último
        dup_mask = adm.duplicated("CODIGO_PRESTACION", keep=False)
        if dup_mask.any():
            st.warning(
                f"{adm.loc[dup_mask, 'CODIGO_PRESTACION'].nunique()} códigos "
                "duplicados en administración (se usa el último valor)."
            )
            with st.expander("Ver duplicados"):
                st.dataframe(adm[dup_mask], use_container_width=True, hide_index=True)
        adm = adm.drop_duplicates("CODIGO_PRESTACION", keep="last")

        # Último valor histórico vigente hasta el período elegido
        # (las valorizaciones no son todos los meses).
        hasta = df[df["PERIODO"] <= compare_period].sort_values("PERIODO")

        # N   = última valorización hasta el período elegido (el "actual").
        # N-1 = última valorización ANTERIOR al período elegido
        #       (no siempre es el mes calendario anterior: hay meses sin valorización).
        vigente = (
            hasta.drop_duplicates("CODIGO_PRESTACION", keep="last")
            [["CODIGO_PRESTACION", "PERIODO", "VALOR"]]
            .rename(columns={"PERIODO": "PERIODO_HISTORICO",
                             "VALOR": "VALOR_HISTORICO"})
        )
        previo = (
            df[df["PERIODO"] < compare_period]
            .sort_values("PERIODO")
            .drop_duplicates("CODIGO_PRESTACION", keep="last")
            [["CODIGO_PRESTACION", "PERIODO", "VALOR"]]
            .rename(columns={"PERIODO": "PERIODO_N1",
                             "VALOR": "VALOR_HIST_N1"})
        )
        cantidad = (
            hasta.groupby("CODIGO_PRESTACION").size().rename("N_VALORIZACIONES")
        )

        t = adm.merge(vigente, on="CODIGO_PRESTACION", how="left")
        t = t.merge(previo, on="CODIGO_PRESTACION", how="left")
        t = t.merge(cantidad, on="CODIGO_PRESTACION", how="left")
        t["N_VALORIZACIONES"] = t["N_VALORIZACIONES"].fillna(0).astype(int)

        existe_algun_dia = set(df["CODIGO_PRESTACION"])
        t["ESTADO"] = "Con valor histórico"
        sin_valor = t["VALOR_HISTORICO"].isna()
        t.loc[sin_valor, "ESTADO"] = "Histórico solo posterior al período"
        t.loc[
            sin_valor & ~t["CODIGO_PRESTACION"].isin(existe_algun_dia), "ESTADO"
        ] = "No existe en histórico"

        # ---- Alfabeta (tercer origen): se suma como columna para comparar ----
        if alfabeta_df is not None:
            alf = alfabeta_df.drop_duplicates("CODIGO_PRESTACION", keep="last")
            alf = alf[["CODIGO_PRESTACION", "VALOR"]].rename(
                columns={"VALOR": "VALOR_ALFABETA"}
            )
            t = t.merge(alf, on="CODIGO_PRESTACION", how="left")
        else:
            t["VALOR_ALFABETA"] = float("nan")

        t["DIFERENCIA"] = t["VALOR"] - t["VALOR_HISTORICO"]
        t["DIFERENCIA_PCT"] = (
            t["DIFERENCIA"] / t["VALOR_HISTORICO"].replace(0, float("nan"))
        ) * 100
        t = t.rename(columns={"VALOR": "VALOR_ADMINISTRACION"})

        base_n1 = t["VALOR_HIST_N1"].replace(0, float("nan"))
        t["DIF_ALFABETA_VS_ADMIN"] = t["VALOR_ALFABETA"] - t["VALOR_ADMINISTRACION"]
        t["DIF_ALFABETA_VS_ADMIN_PCT"] = (
            t["DIF_ALFABETA_VS_ADMIN"]
            / t["VALOR_ADMINISTRACION"].replace(0, float("nan"))
        ) * 100
        t["DIFERENCIA_N1"] = t["VALOR_ADMINISTRACION"] - t["VALOR_HIST_N1"]
        t["DIFERENCIA_PCT_N1"] = t["DIFERENCIA_N1"] / base_n1 * 100
        t["VALORIZADO_EN_PERIODO"] = (
            t["PERIODO_HISTORICO"] == compare_period
        ).map({True: "Sí", False: "No"})

        # Variación ya cargada en el histórico: solo si N pertenece al período actual
        t["VAR_HIST_N_VS_N1_PCT"] = (t["VALOR_HISTORICO"] / base_n1 - 1) * 100
        t.loc[t["VALORIZADO_EN_PERIODO"] == "No", "VAR_HIST_N_VS_N1_PCT"] = float("nan")

        # ---- filtros rápidos ----
        c1, c2, c3, c4, c5 = st.columns([2, 1.2, 1, 1, 1])
        txt = c1.text_input("Buscar en descripción / código", key="adm_txt").strip().lower()
        verde = c3.number_input("Verde hasta ±%", value=2.0, step=0.5)
        amar = c4.number_input("Amarillo hasta ±%", value=10.0, step=1.0)
        tol = c5.number_input(
            "Coincide si dif. ≤ (%)", value=0.1, step=0.1,
            help="Tolerancia para considerar que administración "
                 "es igual a N o a N-1."
        )

        t["SEMAFORO"] = t["DIFERENCIA_PCT"].apply(lambda p: semaforo(p, verde, amar))

        opciones = ["🟢 OK", "🟡 Revisar", "🔴 Desvío", "⚪ Sin dato"]
        sel = c2.multiselect("Semáforo", opciones, default=opciones)

        def _coincide(r):
            if pd.isna(r["VALOR_HISTORICO"]):
                return "Sin histórico"
            if pd.notna(r["DIFERENCIA_PCT"]) and abs(r["DIFERENCIA_PCT"]) <= tol:
                return "Igual a N (ya cargado)"
            if pd.notna(r["DIFERENCIA_PCT_N1"]) and abs(r["DIFERENCIA_PCT_N1"]) <= tol:
                return "Igual a N-1 (valor anterior)"
            return "Distinto de N y N-1"

        t["COINCIDE_CON"] = t.apply(_coincide, axis=1)

        coin_opts = ["Igual a N (ya cargado)", "Igual a N-1 (valor anterior)",
                     "Distinto de N y N-1", "Sin histórico"]
        sel_coin = st.multiselect("Valor de administración coincide con", coin_opts,
                                  default=coin_opts)

        v = t[t["SEMAFORO"].isin(sel) & t["COINCIDE_CON"].isin(sel_coin)]
        if txt:
            v = v[
                v["DESCRIPCION"].astype(str).str.lower().str.contains(txt, regex=False)
                | v["CODIGO_PRESTACION"].str.lower().str.contains(txt, regex=False)
            ]

        # ---- resumen ----
        r = t["SEMAFORO"].value_counts().reindex(opciones, fill_value=0)
        m1, m2, m3, m4 = st.columns(4)
        for col, op in zip((m1, m2, m3, m4), opciones):
            col.metric(op, int(r[op]))

        rc = t["COINCIDE_CON"].value_counts().reindex(coin_opts, fill_value=0)
        n1, n2, n3, n4 = st.columns(4)
        for col, op in zip((n1, n2, n3, n4), coin_opts):
            col.metric(op, int(rc[op]))

        v = v.sort_values("DIFERENCIA_PCT", key=abs, ascending=False)

        mostrar = v[[
            "SEMAFORO", "COINCIDE_CON", "ESTADO",
            "CODIGO_PRESTACION", "DESCRIPCION",
            "VALOR_ADMINISTRACION", "VALOR_ALFABETA",
            "DIF_ALFABETA_VS_ADMIN", "DIF_ALFABETA_VS_ADMIN_PCT",
            "VALOR_HISTORICO", "PERIODO_HISTORICO", "VALORIZADO_EN_PERIODO",
            "DIFERENCIA", "DIFERENCIA_PCT",
            "VALOR_HIST_N1", "PERIODO_N1", "DIFERENCIA_N1", "DIFERENCIA_PCT_N1",
            "VAR_HIST_N_VS_N1_PCT", "N_VALORIZACIONES",
        ]]

        # Selección de fila -> movimiento histórico del código
        evento = st.dataframe(
            mostrar, use_container_width=True, hide_index=True, height=450,
            on_select="rerun", selection_mode="single-row",
            column_config={
                "VALOR_ADMINISTRACION": st.column_config.NumberColumn("Valor admin.", format="$ %.2f"),
                "VALOR_ALFABETA": st.column_config.NumberColumn("Valor Alfabeta", format="$ %.2f"),
                "DIF_ALFABETA_VS_ADMIN": st.column_config.NumberColumn("Dif. Alfabeta vs admin.", format="$ %.2f"),
                "DIF_ALFABETA_VS_ADMIN_PCT": st.column_config.NumberColumn("Dif. % Alfabeta vs admin.", format="%.2f %%"),
                "VALOR_HISTORICO": st.column_config.NumberColumn("Hist. N (último)", format="$ %.2f"),
                "PERIODO_HISTORICO": st.column_config.TextColumn("Período N"),
                "VALORIZADO_EN_PERIODO": st.column_config.TextColumn("¿Valorizado en el período?"),
                "VALOR_HIST_N1": st.column_config.NumberColumn("Hist. N-1", format="$ %.2f"),
                "PERIODO_N1": st.column_config.TextColumn("Período N-1"),
                "DIFERENCIA_N1": st.column_config.NumberColumn("Dif. vs N-1", format="$ %.2f"),
                "DIFERENCIA_PCT_N1": st.column_config.NumberColumn("Dif. % vs N-1", format="%.2f %%"),
                "VAR_HIST_N_VS_N1_PCT": st.column_config.NumberColumn("Var. hist. N vs N-1 %", format="%.2f %%"),
                "DIFERENCIA": st.column_config.NumberColumn("Diferencia", format="$ %.2f"),
                "DIFERENCIA_PCT": st.column_config.NumberColumn("Dif. %", format="%.2f %%"),
                "N_VALORIZACIONES": st.column_config.NumberColumn("Valorizaciones"),
            },
        )
        st.caption(
            f"{len(mostrar)} de {len(t)} códigos. N = última valorización histórica "
            f"vigente hasta {period_label(compare_period)}; N-1 = última valorización "
            f"anterior a {period_label(compare_period)}. "
            "Tocá una fila para ver su movimiento histórico."
        )

        filas_sel = evento.selection.rows if evento is not None else []
        if filas_sel:
            fila = mostrar.iloc[filas_sel[0]]
            cod = fila["CODIGO_PRESTACION"]
            hist = historical_df[historical_df["CODIGO_PRESTACION"] == cod] \
                .sort_values("FECHA_INICIO_VIGENCIA")

            st.subheader(f"Movimiento histórico - {cod}")
            st.caption(str(fila["DESCRIPCION"]))

            if hist.empty:
                st.info("Este código no existe en el histórico.")
            else:
                g = hist.rename(columns={"VALOR": "Histórico"})
                fig = px.line(
                    g, x="FECHA_INICIO_VIGENCIA", y="Histórico",
                    markers=True, title=f"Histórico vs administración - {cod}",
                )
                fig.add_hline(
                    y=float(fila["VALOR_ADMINISTRACION"]),
                    line_dash="dash",
                    annotation_text="Administración",
                )
                fig.update_layout(
                    xaxis_title="Fecha de vigencia", yaxis_title="Valor", height=380
                )
                st.plotly_chart(fig, use_container_width=True)
                st.dataframe(
                    hist[["FECHA_INICIO_VIGENCIA", "FECHA_FIN_VIGENCIA", "VALOR"]],
                    use_container_width=True, hide_index=True,
                )

        st.divider()

        d1, d2 = st.columns(2)

        d1.download_button(
            "⬇ Informe completo (Excel)",
            build_report(
                t, df, compare_period, verde, amar, tol,
                alfabeta=alfabeta_df,
                nombres={"historico": source_name,
                         "administracion": compare_name or "-",
                         "alfabeta": alfabeta_name or "-"},
            ),
            file_name=f"informe_administracion_{compare_period}.xlsx",
            help="Resumen (prestaciones, administración y Alfabeta), listado de administración, hoja Alfabeta, histórico por código, "
                 "mayores desvíos y códigos sin valor histórico.",
        )

        d2.download_button(
            "⬇ Solo lo filtrado en pantalla (Excel)",
            to_excel_bytes({"Comparacion": mostrar}),
            file_name=f"comparacion_{compare_period}.xlsx",
        )